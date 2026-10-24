"""Strict, atomic validation: never calculate from a partly rejected upload."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import load_workbook

MAX_ROWS = 10_000
MAX_BYTES = 5 * 1024 * 1024
REQUIRED_COLUMNS = [
    "invoice_id", "customer_id", "customer_name", "invoice_date", "due_date",
    "payment_date", "net_amount_sar", "vat_amount_sar", "invoice_amount_sar", "currency",
]
UNSUPPORTED_COLUMNS = {
    "paid_amount", "payment_amount", "payment_amount_sar", "amount_paid_sar",
    "credit_note", "credit_note_amount", "credit_note_amount_sar", "invoice_status",
}


@dataclass
class ValidationResult:
    data: pd.DataFrame | None
    errors: pd.DataFrame
    warnings: list[str]

    @property
    def valid(self) -> bool:
        return self.data is not None and self.errors.empty


def read_invoices(contents: bytes, filename: str) -> pd.DataFrame:
    """Load one bounded CSV or first XLSX sheet. Nothing is written to disk."""
    if len(contents) > MAX_BYTES:
        raise ValueError("The file exceeds the 5 MB prototype limit.")
    if filename.lower().endswith(".csv"):
        decoded = contents.decode("utf-8-sig")
        header = next(csv.reader(io.StringIO(decoded)), [])
        if len(header) != len(set(header)):
            raise ValueError("Duplicate column headers are not supported.")
        raw = pd.read_csv(io.StringIO(decoded), dtype=str, keep_default_na=False, nrows=MAX_ROWS + 1)
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError("The workbook exceeds the 25 MB expanded-size limit.")
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        try:
            header = list(next(workbook.worksheets[0].iter_rows(values_only=True), ()))
            if len(header) != len(set(header)):
                raise ValueError("Duplicate column headers are not supported.")
        finally:
            workbook.close()
        raw = pd.read_excel(io.BytesIO(contents), dtype=str, keep_default_na=False, nrows=MAX_ROWS + 1)
    else:
        raise ValueError("Use a UTF-8 CSV or an XLSX workbook; only its first sheet is read.")
    if len(raw) > MAX_ROWS:
        raise ValueError("The file exceeds the 10,000-row prototype limit.")
    return raw


def _parse_day(value: object) -> pd.Timestamp | None:
    if isinstance(value, (datetime, date, pd.Timestamp)):
        result = pd.Timestamp(value)
        return result.normalize() if result == result.normalize() and result.tzinfo is None else None
    text = str(value).strip()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return None
    try:
        parsed = date.fromisoformat(text)
        return pd.Timestamp(parsed) if 1900 <= parsed.year <= 2100 else None
    except (ValueError, OverflowError):
        return None


def _parse_cents(value: object, allow_zero: bool = False) -> int | None:
    text = str(value).strip()
    if not re.fullmatch(r"\d+(?:\.\d{1,2})?", text):
        return None
    try:
        amount = Decimal(text)
        if not amount.is_finite() or amount > Decimal("1000000000"):
            return None
        if amount < 0 or (not allow_zero and amount == 0):
            return None
        return int(amount * 100)
    except (InvalidOperation, ValueError, OverflowError):
        return None


def validate_invoices(raw: pd.DataFrame, as_of: date) -> ValidationResult:
    issues: list[dict] = []
    warnings: list[str] = []

    def issue(row: int | None, field: str, message: str) -> None:
        issues.append({"row": row, "field": field, "message": message})

    missing = sorted(set(REQUIRED_COLUMNS) - set(raw.columns))
    if missing:
        issue(None, "columns", "Missing columns: " + ", ".join(missing))
    unsupported = sorted(UNSUPPORTED_COLUMNS & set(raw.columns))
    if unsupported:
        issue(None, "columns", "Unsupported payment/status/credit-note fields: " + ", ".join(unsupported))
    if raw.columns.duplicated().any():
        issue(None, "columns", "Column headers must be unique.")
    if raw.empty:
        issue(None, "file", "At least one invoice is required.")
    if len(raw) > MAX_ROWS:
        issue(None, "file", "At most 10,000 invoices are supported.")
    if issues:
        return ValidationResult(None, pd.DataFrame(issues), warnings)

    frame = raw.copy().reset_index(drop=True)
    cutoff = pd.Timestamp(as_of)
    start = (cutoff.to_period("M") - 11).start_time
    for name in ["invoice_id", "customer_id", "customer_name"]:
        frame[name] = frame[name].fillna("").astype(str).str.strip()
        for index in frame.index[frame[name].eq("") | frame[name].str.len().gt(160)]:
            issue(int(index) + 2, name, "Must contain 1-160 characters.")
    for index in frame.index[frame["invoice_id"].duplicated(keep=False)]:
        issue(int(index) + 2, "invoice_id", "Duplicate invoice ID; resolve it before analysis.")
    for cid, group in frame.groupby("customer_id"):
        if group["customer_name"].nunique() > 1:
            issue(None, "customer_name", f"Customer ID {cid!r} maps to multiple names.")

    for field in ["invoice_date", "due_date", "payment_date"]:
        parsed = []
        for index, value in frame[field].items():
            blank = pd.isna(value) or str(value).strip() == ""
            if blank and field == "payment_date":
                parsed.append(pd.NaT)
                continue
            day = _parse_day(value)
            if day is None:
                issue(int(index) + 2, field, "Use a valid ISO date (YYYY-MM-DD), within 1900-2100.")
            parsed.append(day if day is not None else pd.NaT)
        frame[field] = pd.to_datetime(parsed)

    for source, target, allow_zero in [
        ("net_amount_sar", "net_cents", False),
        ("vat_amount_sar", "vat_cents", True),
        ("invoice_amount_sar", "gross_cents", False),
    ]:
        parsed = []
        for index, value in frame[source].items():
            cents = _parse_cents(value, allow_zero)
            if cents is None:
                issue(int(index) + 2, source, "Use a plain nonnegative amount with at most 2 decimals (positive for net/gross), maximum SAR 1 billion.")
            parsed.append(cents if cents is not None else 0)
        frame[target] = pd.Series(parsed, dtype="int64")
    for index, row in frame.iterrows():
        line = int(index) + 2
        if str(row["currency"]).strip() != "SAR":
            issue(line, "currency", "Only SAR is supported; no currency conversion is performed.")
        if row["net_cents"] + row["vat_cents"] != row["gross_cents"]:
            issue(line, "invoice_amount_sar", "Gross must equal net plus VAT exactly.")
        if pd.notna(row["invoice_date"]) and not start <= row["invoice_date"] <= cutoff:
            issue(line, "invoice_date", "Invoice is outside the selected 12-month review window.")
        if row["due_date"] < row["invoice_date"]:
            issue(line, "due_date", "Due date cannot precede issue date.")
        if row["payment_date"] < row["invoice_date"]:
            issue(line, "payment_date", "Payment cannot precede issue date.")
        if row["payment_date"] > cutoff:
            issue(line, "payment_date", "Payment is after the snapshot; provide an as-of snapshot instead.")
        if row["due_date"] > row["invoice_date"] + pd.Timedelta(days=365):
            warnings.append(f"Row {line}: terms exceed 365 days; confirm this is intentional.")

    months = frame["invoice_date"].dropna().dt.to_period("M").nunique()
    if months < 12:
        warnings.append(f"Only {months} of 12 months contain invoices. Missing months may mean no sales or incomplete files; confirm with the applicant.")
    extra = sorted(set(raw.columns) - set(REQUIRED_COLUMNS) - {"category"})
    if extra:
        warnings.append("Extra columns are not analyzed: " + ", ".join(map(str, extra)))
    if issues:
        return ValidationResult(None, pd.DataFrame(issues), warnings)

    frame["is_paid"] = frame["payment_date"].notna()
    frame["days_past_due"] = (cutoff - frame["due_date"]).dt.days.clip(lower=0)
    frame.loc[frame["is_paid"], "days_past_due"] = 0
    frame["status"] = "Open - within terms"
    frame.loc[~frame["is_paid"] & frame["due_date"].lt(cutoff), "status"] = "Overdue"
    frame.loc[frame["is_paid"], "status"] = "Paid"
    for column in ["net_amount_sar", "vat_amount_sar", "invoice_amount_sar"]:
        frame[column] = pd.to_numeric(frame[column])
    return ValidationResult(frame, pd.DataFrame(columns=["row", "field", "message"]), warnings)


def safe_csv(frame: pd.DataFrame) -> bytes:
    """Avoid spreadsheet-formula execution in exported user-provided text."""
    def escape(value: object) -> object:
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    return frame.map(escape).to_csv(index=False).encode("utf-8-sig")
